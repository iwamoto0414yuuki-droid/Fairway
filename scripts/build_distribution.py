"""Build a recipient-ready Windows guide and a self-contained Fairway ZIP.
Requires python-docx, reportlab and the native screenshots in release/.
No personal account, certificate, pairing file or round data is distributed.
"""
from pathlib import Path
import argparse
import hashlib
import json
import shutil
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image

ROOT = Path(__file__).resolve().parents[1]
RELEASE = ROOT / 'release'
OUT = RELEASE / '配布用_Fairway_0.2.0_Windows'
IPA_NAME = 'Fairway-0.2.0-full.ipa'
GUIDE_NAME = '02_導入手順_Windows'
GUIDE_BANNER = 'FAIRWAY 0.2.0  /  Windows 導入ガイド'
CHECKED_DATE = '2026-10-24'
EXPECTED = '5a9786d519dc59ed9adaccbdc65e982b787b69c99ee635aa9b71ebdade07ec66'
SOURCES = [
    ('S1', 'SideStore：事前準備（Windows条件・必要ソフト）', 'https://docs.sidestore.io/docs/installation/prerequisites'),
    ('S2', 'SideStore：初回インストール', 'https://docs.sidestore.io/docs/installation/install'),
    ('S3', 'SideStore：期限・アプリ枠・初回Refresh', 'https://docs.sidestore.io/docs/faq'),
    ('S4', 'iloader公式サイト（Windows版ダウンロード）', 'https://iloader.app/'),
    ('S5', 'LocalDevVPN（App Store）', 'https://apps.apple.com/app/localdevvpn/id6755608044'),
    ('S6', 'Apple：WindowsでiPhoneを認識しない場合', 'https://support.apple.com/ja-jp/108643'),
    ('S7', 'SideStore：ペアリング情報の入れ直し', 'https://docs.sidestore.io/docs/advanced/pairing-file'),
]
URLS = {key: url for key, _, url in SOURCES}

# All paragraphs are shared by the PDF and Word versions.
# Each page has a deliberately bounded amount of content for A4 printing.
PAGES = [
    ('受け取った人へ', [
        ('lead', 'Fairway 0.2.0｜WindowsからiPhone・Apple Watchへ'),
        ('box', '入れるIPAは1本だけです。最新版には、最初のアプリ本体と追加機能がすべて含まれます。旧版を先に入れたり、2本を合体したりする操作は不要です。'),
        ('h', '配布セットの中身'),
        ('p', 'Fairway-0.2.0-full.ipa：初回インストール・既存アプリ更新の両方に使う本体。iPhone用とWatch用を同梱しています。'),
        ('p', '02_導入手順_Windows.pdf：この説明書。Word版は追記・編集用です。'),
        ('h', '用意するもの'),
        ('p', 'iOS 17以上のiPhone。Watchで使う場合はwatchOS 10以上の、iPhoneとペアリング済みApple Watchも用意します。'),
        ('p', 'Windowsパソコン、データ通信対応USBケーブル、Wi-Fi、受け取った本人のApple Accountと確認コードを受け取れる環境。'),
        ('p', 'Windowsは64bit版を使用します。公式手順では32bit版とWindows 10 on Armは対象外です。[S1]'),
        ('h', '最初に知っておくこと'),
        ('p', 'アプリの作成者のApple Accountは使いません。利用する本人が自分のアカウントでiloaderとSideStoreにログインします。パスワード・確認コード・端末のペアリング情報を作成者へ送る必要はありません。'),
        ('p', 'SideStore等は公式リンクから取得します。このZIPには他社のインストーラーや署名情報を入れていません。'),
        ('p', 'Fairwayのネイティブテスト30件とビルドは成功しています。ただし通常のSideStore経由のWatch導入には、整合性エラーの実機報告があり未解決です。Watchを使う場合は、別の署名・導入経路の検討が必要です。'),
        ('p', '流れ：Windowsの準備 → SideStore導入 → Fairway導入 → Watch確認 → 普段の期限更新。'),
    ]),
    ('WindowsでSideStoreを入れる', [
        ('lead', 'パソコンを使うのは初回設定です'),
        ('h', '1　受け取ったZIPを展開する'),
        ('p', 'ZIPを右クリックして「すべて展開」。展開したフォルダ内にFairway-0.2.0-full.ipaがあることを確認します。IPA自体は展開しません。'),
        ('h', '2　iTunesとiloaderを準備する'),
        ('p', '公式の事前準備ページ[S1]のWindows欄からiTunesを導入します。iloader公式サイト[S4]のWindows欄でMSI Installerを選び、インストールします。'),
        ('link', ('iloader公式ダウンロードを開く', URLS['S4'])),
        ('p', 'ブラウザー内で動く試験版ではなく、Windowsに入れる通常版を使います。画面や配布名が変わった場合は公式ページのWindows用案内を優先してください。'),
        ('h', '3　iPhoneをUSBでつなぐ'),
        ('p', 'iPhoneのロックを解除してUSB接続します。iPhoneに確認が出たら、このパソコンを「信頼」します。認識しない場合は7ページを参照。[S6]'),
        ('h', '4　iloaderからSideStoreを入れる'),
        ('p', 'iloaderを開き、本人のApple Accountでログイン。対象のiPhoneを選び、「Install SideStore (Stable)」を実行します。[S2]'),
        ('box', 'ここで入れるのはSideStoreです。Fairwayは、iPhoneでSideStoreの初期設定を終えてから入れます。'),
        ('p', '完了したらiPhoneのホーム画面にSideStoreがあるか確認し、次のページへ進みます。'),
    ]),
    ('iPhoneでSideStoreを使えるようにする', [
        ('lead', '初回だけ行う設定'),
        ('h', '1　アプリの開発元を信頼する'),
        ('p', '設定 → 一般 → VPNとデバイス管理を開き、開発用App欄の自分のApple Accountを選んで信頼します。新しいiOSでは「許可して再起動」等が表示されるので、その案内に従います。[S2]'),
        ('h', '2　デベロッパモードを有効にする'),
        ('p', '設定 → プライバシーとセキュリティ → デベロッパモードをオン。再起動後の確認も完了します。設定項目がない場合は、SideStoreの導入と信頼が済んでいるか確認します。[S2]'),
        ('h', '3　LocalDevVPNを入れて接続する'),
        ('p', 'iPhoneのApp StoreからLocalDevVPNを入れ、Connectで接続。VPN構成追加の確認には端末上で許可します。SideStoreでインストール・更新・期限延長する時に必要です。[S1・S5]'),
        ('link', ('LocalDevVPNをApp Storeで開く', URLS['S5'])),
        ('h', '4　SideStore自身を一度Refreshする'),
        ('p', 'Wi-Fiに接続してSideStoreを開き、iloaderと同じApple Accountでログイン。My AppsでSideStore横の残日数表示を押してRefreshします。完了してからFairwayへ進みます。[S2・S3]'),
        ('box', '成功の目安：SideStoreが起動し、自分自身の期限更新が成功すること。まだFairwayが入っていなくても正常です。'),
        ('p', '署名証明書の作成・失効を確認された場合、他のサイドロードアプリがある人は影響する対象を確認してください。初回設定で分からない場合は、表示文を控えてから相談します。'),
    ]),
    ('FairwayとWatchを入れる', [
        ('lead', '初回も最新版IPAだけを使います'),
        ('h', '1　IPAをiPhoneの「ファイル」に保存する'),
        ('p', '展開済みのFairway-0.2.0-full.ipaをOneDrive・iCloud Drive等でiPhoneへ渡し、「ファイル」から選べる状態にします。ZIPをiPhoneへ渡した場合は「ファイル」でZIPをタップして展開します。'),
        ('h', '2　SideStoreでIPAを選ぶ'),
        ('p', 'Wi-FiとLocalDevVPNに接続。SideStore → My Apps → ＋からIPAを選びます。完了するとホーム画面にFairwayが追加されます。IPAをタップするだけではインストールされません。'),
        ('p', 'Watchや付属機能を除去する確認が出た場合、Watchを残す必要があります。選択内容が分からなければ、その画面を控えて確認してください。'),
        ('p', 'InvalidCompanionAppBundleIdentifierが出た場合は、エラー内で要求されたiPhoneアプリIDに合わせたIPAが必要です。そのエラー文を作成者へ伝え、修正版を受け取って同じ手順で再試行します。SideStore初期設定のやり直しは不要です。'),
        ('h', '3　Watchへの導入を確認する'),
        ('p', 'iPhoneの「Watch」アプリを開き、Fairwayが表示されたらインストールします。iPhone・Watchの両方でFairwayを起動してください。Watch側で開発者モード等の案内が出た場合は、端末の案内に従います。'),
        ('box', 'Watch用アプリは同梱済みですが、通常のSideStoreでiPhoneに入った後、Watchの整合性検証で止まる実機報告があります。ID修正版だけではWatch用署名の問題は解決しません。止まった場合は再インストールを繰り返さず、作成者へ伝えてください。'),
        ('h', '4　スコアが届くか確認する'),
        ('p', 'iPhoneでラウンドを作成し、Watchに届くまで待ちます。届かない時は、両端末でFairwayを開き、iPhoneのラウンド内の「このラウンドを再送」を押します。'),
        ('p', 'Watchで1ホールを保存し、iPhoneのカードへ反映されるか確認します。転送は即時に届くとは限りません。'),
    ]),
    ('Fairwayの使い方', [
        ('h', 'ラウンド前：PARとグリーンを登録'),
        ('p', 'iPhoneで9H／18Hを作成。各ホールの「PAR・グリーン位置を設定」を開き、ゴルフ場検索 → 航空写真を拡大 → グリーン中央をタップして保存します。初期PARは全て4なので実際の値へ変更してください。'),
        ('h', 'プレー中：打数・距離・クラブ'),
        ('p', 'Watchで打数を−／＋し「保存して次へ」。パット数は任意です。罰打は打数に含めます。クラブは「使用クラブ（任意）」で一打ごとに選び、戻ってスコアを保存します。'),
        ('p', '距離は「GPS距離」→「距離を更新」で取得。グリーン中央までの直線距離をヤードで表示します。測定時刻と精度を確認してください。歩いても距離は自動更新されません。'),
        ('images', ['native-watch.png', 'native-watch-summary.png', 'native-watch-clubs.png']),
        ('p', '左：打数入力　中央：集計　右：任意のクラブ記録（シミュレーター画像）'),
        ('h', 'スコアを見る・次回へ引き継ぐ'),
        ('p', 'ラウンド画面で前半・後半・全体の打数とパット合計を確認できます。パット未入力は0と区別して記録済みホール数を表示します。9Hは全体のみです。'),
        ('p', '終了後はカードを確認してラウンドを終了。同じコースは、履歴の「同じコースで新規ラウンド」からPARとグリーン位置を再利用できます。'),
    ]),
    ('期限のRefreshと、アプリの更新', [
        ('lead', '「期限の更新」と「新しい版を入れる」は別の操作です'),
        ('h', '毎週：使える期限を延ばす'),
        ('p', '無料Apple Accountでは通常7日ごとのRefreshが必要です。Wi-FiとLocalDevVPNに接続し、SideStoreのMy AppsでSideStoreとFairwayの残日数を確認して更新します。[S3]'),
        ('p', '残り日数が少なくなる前に確認する習慣をつけてください。自動更新だけに頼らず、ラウンド前に起動と期限を確認します。'),
        ('h', '新しいIPAを受け取った時：機能を更新する'),
        ('p', '今のFairwayを削除せず、新しいIPAを「ファイル」に保存し、同じSideStore・同じApple Accountで＋から選んで上書きします。iPhoneとWatchの両方を更新してから使います。'),
        ('p', '今回の0.2.0には初版本体、前半・後半／パット集計、クラブ記録が含まれています。初回・更新でIPAを使い分ける必要はありません。'),
        ('h', '電池を抑える使い方'),
        ('p', 'GPSは必要な地点だけで更新します。クラブ選択では測位せず、スコア保存時にまとめて同期します。長時間の電池消費は実測していません。'),
        ('p', 'LocalDevVPNはSideStoreで導入・Refreshする時に接続します。Fairwayのスコア入力やGPS距離計算自体に、そのVPNは不要です。'),
        ('box', '期限切れや更新失敗の時も、保存データを残すため、まずアプリを削除せず状態を確認してください。'),
        ('p', '無料アカウントのSideStoreには、SideStore自身を含め同時3アプリ等の制限があります。枠に関するメッセージは公式FAQで確認してください。[S3]'),
    ]),
    ('困った時と、公式リンク', [
        ('h', '最初に確認すること'),
        ('p', 'WindowsにiPhoneが出ない：ロック解除、データ通信対応ケーブル、別のUSBポート、信頼の確認を行います。Apple公式の認識トラブル案内[S6]も参照します。'),
        ('p', 'SideStoreが更新できない：Wi-Fi、LocalDevVPN、ログインしたApple Account、SideStore自身のRefreshを確認します。端末更新後のペアリング問題は公式の入れ直し案内[S7]を参照します。'),
        ('p', 'WatchにFairwayがない：Watchアプリのインストール候補とエラーを確認。iPhoneに入っただけでWatchも成功したとは扱わず、表示文と端末情報を作成者へ伝えてください。'),
        ('p', 'GPS距離が出ない：対象ホールのグリーン登録、位置情報許可、屋外の受信状態を確認し、距離を更新。測位に失敗しても打数は記録できます。'),
        ('h', '相談する時に伝える情報'),
        ('p', 'どの手順で止まったか、エラー文、iPhoneの機種／iOS、Watchの機種／watchOS、SideStoreとiloaderの版。Apple Accountのパスワードや確認コードは送りません。'),
        ('h', '公式リンク（2026年10月24日確認）'),
        *[('link', (f'[{key}] {title}', url)) for key, title, url in SOURCES],
        ('p', '画面名・導入条件が変わった場合は公式手順を確認してください。本書のFairway操作は0.2.0を対象にしています。'),
    ]),
]


def hyperlink(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '276343'); prop.append(color)
    run.append(prop); text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def pdf():
    pdfmetrics.registerFont(TTFont('MeiryoGuide', 'C:/Windows/Fonts/meiryo.ttc', subfontIndex=0))
    pdfmetrics.registerFont(TTFont('MeiryoGuideBold', 'C:/Windows/Fonts/meiryob.ttc', subfontIndex=0))
    styles = {
        'title': ParagraphStyle('title', fontName='MeiryoGuideBold', fontSize=21, leading=29, textColor=colors.HexColor('#153f2b'), spaceAfter=14, wordWrap='CJK'),
        'lead': ParagraphStyle('lead', fontName='MeiryoGuideBold', fontSize=11, leading=17, textColor=colors.HexColor('#276343'), spaceAfter=12, wordWrap='CJK'),
        'h': ParagraphStyle('h', fontName='MeiryoGuideBold', fontSize=12, leading=18, textColor=colors.HexColor('#276343'), spaceBefore=10, spaceAfter=5, wordWrap='CJK'),
        'p': ParagraphStyle('p', fontName='MeiryoGuide', fontSize=10, leading=16, spaceAfter=7, wordWrap='CJK'),
        'link': ParagraphStyle('link', fontName='MeiryoGuide', fontSize=9, leading=15, textColor=colors.HexColor('#276343'), spaceAfter=6, wordWrap='CJK'),
        'box': ParagraphStyle('box', fontName='MeiryoGuideBold', fontSize=10, leading=16, borderPadding=9, backColor=colors.HexColor('#edf4eb'), spaceBefore=8, spaceAfter=13, wordWrap='CJK'),
    }
    flow = []
    for page, (title, blocks) in enumerate(PAGES):
        if page: flow.append(PageBreak())
        flow.append(Paragraph(f'{page + 1:02d}　{escape(title)}', styles['title']))
        for kind, content in blocks:
            if kind == 'images':
                row = [Image(str(RELEASE / name), width=118, height=118 * 394 / 324) for name in content]
                table = Table([row], colWidths=[165] * 3)
                table.setStyle(TableStyle([('ALIGN', (0, 0), (-1, -1), 'CENTER'), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
                flow.extend([Spacer(1, 5), table, Spacer(1, 7)])
            elif kind == 'link':
                label, url = content
                flow.append(Paragraph(f'<link href="{escape(url, {chr(34): "&quot;"})}">{escape(label)}</link>', styles[kind]))
            else: flow.append(Paragraph(escape(content), styles[kind]))
    def frame(canvas, document):
        canvas.setFont('MeiryoGuide', 8)
        canvas.setFillColor(colors.HexColor('#68766b'))
        canvas.drawString(48, 817, GUIDE_BANNER)
        canvas.drawRightString(547, 25, f'{document.page} / {len(PAGES)}')
        canvas.drawString(48, 25, '確認日 2026-10-24')
    doc = SimpleDocTemplate(str(OUT / f'{GUIDE_NAME}.pdf'), pagesize=A4, leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=44, title='Fairway 導入手順 Windows', author='Fairway')
    doc.build(flow, onFirstPage=frame, onLaterPages=frame)


def word():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height=Cm(29.7); sec.page_width=Cm(21)
    sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style=doc.styles[name]; style.font.name='Meiryo'; style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Meiryo')
    normal=doc.styles['Normal']; normal.font.size=Pt(10)
    normal.paragraph_format.line_spacing=1.12; normal.paragraph_format.space_after=Pt(5)
    for name in ['Title','Heading 1','Heading 2']: doc.styles[name].font.color.rgb=RGBColor.from_string('276343')
    doc.styles['Title'].font.size=Pt(21); doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
    sec.header.paragraphs[0].text=GUIDE_BANNER
    footer=sec.footer.paragraphs[0]; footer.text=f'確認日 {CHECKED_DATE}　 /　'
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.core_properties.title='Fairway 導入手順 Windows'; doc.core_properties.author='Fairway'
    for page, (title, blocks) in enumerate(PAGES):
        if page: doc.add_page_break()
        doc.add_paragraph(f'{page+1:02d}　{title}', style='Title')
        for kind, content in blocks:
            if kind=='images':
                table=doc.add_table(rows=1,cols=3)
                for cell,name in zip(table.rows[0].cells,content):
                    cell.paragraphs[0].add_run().add_picture(str(RELEASE/name),width=Cm(4.1))
            elif kind=='link': hyperlink(doc.add_paragraph(),*content)
            elif kind=='h': doc.add_paragraph(content,style='Heading 2')
            else:
                paragraph=doc.add_paragraph(content)
                if kind in ['box','lead']:
                    for run in paragraph.runs: run.bold=True
                if kind=='box':
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EDF4EB'); paragraph._p.get_or_add_pPr().append(shade)
    doc.save(OUT/f'{GUIDE_NAME}.docx')


def build(team_id=None):
    global OUT, IPA_NAME, EXPECTED
    source=RELEASE/'Fairway-0.2.0-unsigned.ipa'
    assert hashlib.sha256(source.read_bytes()).hexdigest()==EXPECTED
    verification=RELEASE/'verification.json'
    suffix=''
    account_note=''
    if team_id:
        from build_sidestore_ipa import prepare
        prepare(team_id)
        suffix='_'+team_id
        source=RELEASE/('SideStore_'+team_id)/('Fairway-0.2.0-SideStore-'+team_id+'.ipa')
        verification=source.parent/'verification.json'
        EXPECTED=hashlib.sha256(source.read_bytes()).hexdigest()
        IPA_NAME=source.name
        OUT=RELEASE/('配布用_Fairway_0.2.0_Windows'+suffix)
        account_note='このIPAはSideStoreがiPhone側IDを local.fairway.score.'+team_id+' にするアカウント専用です。別のApple Accountにはそのまま使わず、エラーに出た必要IDを作成者へ伝えてください。'
        PAGES[0][1].insert(1,('box',account_note))
        for _,blocks in PAGES:
            for index,(kind,content) in enumerate(blocks):
                if isinstance(content,str):blocks[index]=(kind,content.replace('Fairway-0.2.0-full.ipa',IPA_NAME))
    OUT.mkdir(exist_ok=True)
    shutil.copyfile(source,OUT/IPA_NAME)
    pdf(); word()
    # A portable text copy keeps the same substantive instructions and links.
    chunks=['# Fairway 導入手順（Windows）\n\n確認日：2026-10-24\n']
    for page,(title,blocks) in enumerate(PAGES):
        chunks.append(f'\n## {page+1} {title}\n')
        for kind, content in blocks:
            if kind=='images': continue
            if kind=='link': chunks.append(f'[{content[0]}]({content[1]})\n')
            elif kind=='h': chunks.append('### '+content+'\n')
            else: chunks.append(content+'\n')
    (OUT/'02_導入手順_Windows.md').write_text('\n'.join(chunks),encoding='utf-8')
    welcome='''Fairway 0.2.0 配布セット（Windows向け）

初版本体と追加機能は、同梱の最新版IPAにすべて入っています。
初回の人も、このIPAだけをインストールしてください。

1. まず「02_導入手順_Windows.pdf」を読む。
2. Windowsと本人のApple AccountでSideStoreを設定する。
3. Fairway-0.2.0-full.ipaをiPhoneのSideStoreから選ぶ。
4. Watchへの導入とスコアの同期を確認する。

既にFairwayがある場合は削除せず上書き更新します。
PDFとWordは同じ手順です。アプリ本体は1本だけです。
SideStore等の他社インストーラー・署名情報・個人のスコアは含みません。
WatchのSideStore経由での実機導入と電池消費は未検証です。
'''
    (OUT/'01_はじめに.txt').write_text(welcome.replace('Fairway-0.2.0-full.ipa',IPA_NAME)+'\n'+account_note+'\n',encoding='utf-8-sig')
    (OUT/'SHA256.txt').write_text(EXPECTED+'  '+IPA_NAME+'\n',encoding='ascii')
    shutil.copyfile(verification,OUT/'検証記録.json')
    archive=RELEASE/('Fairway-0.2.0-Windows-配布セット'+suffix+'.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for file in sorted(OUT.iterdir()):
            if file.is_file(): z.write(file,OUT.name+'/'+file.name)
    print(json.dumps({'folder':str(OUT),'zip':str(archive),'ipa_sha256':EXPECTED,'zip_bytes':archive.stat().st_size},ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--team-id')
    build(parser.parse_args().team_id)
